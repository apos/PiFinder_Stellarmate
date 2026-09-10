#!/usr/bin/env python3
"""PFSM multi-part presentation. Regenerate: see README.md in this dir."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os, random

_R = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "images"))
_MAP = {
    "hero.jpg":        os.path.join(_R, "readme", "PiFinder.jpg"),
    "cc_full.png":     os.path.join(_R, "readme", "cc_full_page.png"),
    "cc_coupling.png": os.path.join(_R, "readme", "cc_mount_bridge_coupling.png"),
}

DARK  = RGBColor(0x0B, 0x10, 0x26)
DARK2 = RGBColor(0x16, 0x20, 0x40)
LIGHT = RGBColor(0xF7, 0xF9, 0xFC)
NAVY  = RGBColor(0x1B, 0x2A, 0x4A)
BLUE  = RGBColor(0x4E, 0xA8, 0xDE)
AMBER = RGBColor(0xFF, 0xB7, 0x03)
INK   = RGBColor(0x18, 0x20, 0x30)
PAPER = RGBColor(0xEA, 0xEF, 0xF8)
MUTE  = RGBColor(0x5C, 0x67, 0x7B)
MOD   = RGBColor(0x9A, 0xA6, 0xC2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE  = RGBColor(0xD8, 0xE0, 0xEC)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
RED   = RGBColor(0xC6, 0x45, 0x3A)
CARD  = RGBColor(0xFF, 0xFF, 0xFF)
FONT  = "Calibri"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
M = 0.7

# type scale (roughly 2x the previous body sizes)
T_TITLE=40; T_KICK=15; T_H2=26; T_BODY=25; T_BODYSM=21; T_CAP=15
T_SECNUM=96; T_SECT=52; T_SECSUB=22

def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background(); r.shadow.inherit = False
    return s

def _run(r, t, sz, b, c, it=False):
    r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.italic=it; r.font.name=FONT; r.font.color.rgb=c

def text(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.0,sa=4,wrap=True):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,p in enumerate(paras):
        par=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        par.alignment=align; par.line_spacing=ls; par.space_after=Pt(sa); par.space_before=Pt(0)
        for tup in p:
            it=tup[4] if len(tup)>4 else False
            _run(par.add_run(),tup[0],tup[1],tup[2],tup[3],it)
    return tb

def bullets(s,x,y,w,h,items,sz=T_BODY,color=INK,gap=16,marker=BLUE,ls=1.12,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=ls; p.space_after=Pt(gap); p.space_before=Pt(0)
        pPr=p._p.get_or_add_pPr()
        pPr.set('marL',str(int(0.42*914400))); pPr.set('indent',str(-int(0.42*914400)))
        _run(p.add_run(),"▸  ",sz,True,marker)
        if isinstance(it,tuple):
            _run(p.add_run(),it[0],sz,True,color)
            if it[1]: _run(p.add_run(),it[1],sz,False,color)
        else:
            _run(p.add_run(),it,sz,False,color)
    return tb

def circle(s,x,y,d,fill,glyph=None,gc=WHITE,gs=20):
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb=fill; c.line.fill.background(); c.shadow.inherit=False
    if glyph is not None:
        tf=c.text_frame; tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        _run(p.add_run(),glyph,gs,True,gc)
    return c

def pic(s,path,x,y,w,h=None,border=LINE,shadow=True):
    p=s.shapes.add_picture(_MAP[path],Inches(x),Inches(y),Inches(w),Inches(h) if h else None)
    if border is not None: p.line.color.rgb=border; p.line.width=Pt(1)
    if shadow:
        sp=p._element.spPr
        el=sp.makeelement(qn('a:effectLst'),{})
        sh=el.makeelement(qn('a:outerShdw'),{'blurRad':'90000','dist':'34000','dir':'5400000','rotWithShape':'0'})
        clr=sh.makeelement(qn('a:srgbClr'),{'val':'0B1026'})
        clr.append(clr.makeelement(qn('a:alpha'),{'val':'20000'}))
        sh.append(clr); el.append(sh); sp.append(el)
    return p

def card(s,x,y,w,h,fill=CARD,line=LINE):
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    r.adjustments[0]=0.045
    r.fill.solid(); r.fill.fore_color.rgb=fill; r.line.color.rgb=line; r.line.width=Pt(1); r.shadow.inherit=False
    return r

def head(s,kick,title,tc=INK):
    circle(s,M,0.62,0.17,AMBER)
    text(s,M+0.32,0.58,SW-2*M-0.32,0.34,[[(kick.upper(),T_KICK,True,BLUE)]])
    text(s,M,0.98,SW-2*M,1.0,[[(title,T_TITLE,True,tc)]],ls=1.02)

def footer(s,txt,dark=False):
    text(s,M,SH-0.5,SW-2*M,0.3,[[(txt,12,False,MOD if dark else MUTE)]])

def scrim(s,x,y,w,h,color,opacity):
    """opacity 0-100 (0 = invisible)."""
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=color; r.line.fill.background(); r.shadow.inherit=False
    sf=r.fill._xPr.find(qn('a:solidFill')); srgb=sf.find(qn('a:srgbClr'))
    srgb.append(srgb.makeelement(qn('a:alpha'),{'val':str(int(opacity*1000))}))
    return r

def cover(s,path,img_w_px,img_h_px):
    """full-bleed background image, aspect preserved, centered (edges cropped)."""
    ar=img_w_px/img_h_px
    if ar >= SW/SH:
        w=SH*ar; h=SH; x=(SW-w)/2; y=0.0
    else:
        w=SW; h=SW/ar; x=0.0; y=(SH-h)/2
    return s.shapes.add_picture(_MAP[path],Inches(x),Inches(y),Inches(w),Inches(h))

def stars(s,n,seed):
    random.seed(seed)
    for _ in range(n):
        dx=random.uniform(0.15,SW-0.15); dy=random.uniform(0.15,SH-0.15)
        dd=random.choice([0.02,0.03,0.045])
        o=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(dx),Inches(dy),Inches(dd),Inches(dd))
        o.fill.solid(); o.fill.fore_color.rgb=RGBColor(0x3A,0x4A,0x78); o.line.fill.background(); o.shadow.inherit=False

def section(num,kick,title,sub=None):
    s=slide(DARK); stars(s,44,sum(ord(c) for c in num)+len(prs.slides._sldIdLst))
    text(s,M,2.1,3.0,1.7,[[(num,T_SECNUM,True,AMBER)]])
    text(s,M,3.66,SW-2*M,0.4,[[(kick.upper(),T_KICK+1,True,BLUE)]])
    text(s,M,4.06,SW-2*M,1.4,[[(title,T_SECT,True,WHITE)]],ls=1.02)
    if sub: text(s,M,5.3,SW-2*M-1.0,1.2,[[(sub,T_SECSUB,False,MOD)]],ls=1.3)
    return s

# ============================================================

# 1 TITLE — full-bleed hero, aspect preserved, right-aligned (crops the baked logo off the left)
s=slide(DARK)
_hw=SH*(1400/719)
s.shapes.add_picture(_MAP["hero.jpg"],Inches(SW-_hw),0,Inches(_hw),Inches(SH))
scrim(s,0,0,SW,SH,DARK,16)          # gentle overall
scrim(s,0,0,9.6,SH,DARK,24)         # mid band
scrim(s,0,0,7.0,SH,DARK,26)         # strongest behind the text
text(s,M,1.45,7.2,0.4,[[("PiFinder ",18,True,BLUE),("on ",18,True,PAPER),("StellarMate",18,True,AMBER)]])
text(s,M,1.95,7.2,2.3,[[("Plate-solving push-to,",39,True,WHITE)],
                        [("on a full imaging rig.",39,True,WHITE)]],ls=1.12)
text(s,M,4.05,6.6,1.2,[[("One Raspberry Pi. PiFinder answers “where am I pointing”; "
    "StellarMate drives the rest — over INDI.",20,False,PAPER)]],ls=1.35)
text(s,M,5.45,7.0,0.35,[[("github.com/apos/PiFinder_Stellarmate",15,True,BLUE)]])
text(s,M,5.82,7.6,0.3,[[("Community project — not affiliated with PiFinder or StellarMate",12,False,RGBColor(0xC2,0xCB,0xDB))]])

# 2 AGENDA
s=slide(LIGHT); head(s,"Agenda","What this deck covers")
parts=[("1","Project summary & focus"),("2","PFSM Control Center"),
       ("3","PF LX200 & PF Mount Bridge"),("4","PF Simulator")]
for i,(n,t) in enumerate(parts):
    yy=2.2+i*1.02
    circle(s,M,yy,0.62,NAVY,n,WHITE,24)
    text(s,M+0.95,yy+0.03,11.0,0.5,[[(t,T_H2,True,INK)]])
text(s,M,SH-1.9,SW-2*M,1.2,[[("Then two takeaways — ",T_BODYSM,False,MUTE),
    ("PiFinder",T_BODYSM,True,INK),(": where to send PRs.  ",T_BODYSM,False,MUTE),
    ("StellarMate",T_BODYSM,True,INK),(": how the SMOS app could drive PFSM.",T_BODYSM,False,MUTE)]],ls=1.3)

# 3 DUMMIES 1
s=slide(LIGHT); head(s,"For dummies · 1 / 3","Two problems, one telescope")
card(s,M,2.05,6.0,4.2)
circle(s,M+0.4,2.4,0.5,AMBER,"?",DARK,22)
text(s,M+1.1,2.5,4.6,0.5,[[("Where am I\npointing?",T_H2,True,INK)]],ls=1.05)
text(s,M+0.4,3.7,5.2,2.3,[[("A Dobson or push-to EQ has no encoders — the software "
    "has no idea where the tube looks.",T_BODYSM,False,MUTE)]],ls=1.28)
card(s,6.9,2.05,6.0,4.2)
circle(s,7.3,2.4,0.5,BLUE,"★",WHITE,19)
text(s,8.0,2.5,4.6,0.5,[[("Drive the\nwhole rig",T_H2,True,INK)]],ls=1.05)
text(s,7.3,3.7,5.2,2.3,[[("Camera, mount, guiding, focus — all need one coordinated "
    "control layer: INDI, KStars/Ekos, an app.",T_BODYSM,False,MUTE)]],ls=1.28)
footer(s,"For dummies")

# 4 DUMMIES 2
s=slide(LIGHT); head(s,"For dummies · 2 / 3","Each project owns one half")
card(s,M,2.05,6.0,4.2)
text(s,M+0.45,2.35,5.0,0.5,[[("PiFinder",T_H2,True,BLUE)]])
bullets(s,M+0.45,3.15,5.2,3.0,[
    "Global-shutter camera + plate-solver.",
    "Live sky position + push-to arrows.",
    "Python, open hardware & software.",
],sz=T_BODYSM,gap=14)
card(s,6.9,2.05,6.0,4.2)
text(s,7.35,2.35,5.0,0.5,[[("StellarMate",T_H2,True,AMBER)]])
bullets(s,7.35,3.15,5.2,3.0,[
    "INDI — the device-driver standard.",
    "KStars / Ekos — capture & automation.",
    "The StellarMate phone/tablet app.",
],sz=T_BODYSM,gap=14,marker=AMBER)
footer(s,"For dummies")

# 5 DUMMIES 3
s=slide(DARK); stars(s,34,5)
text(s,M,0.62,3.0,0.35,[[("FOR DUMMIES · 3 / 3",T_KICK,True,BLUE)]])
text(s,M,1.0,11.6,1.0,[[("PFSM is the wiring. INDI is the connector.",T_TITLE-2,True,WHITE)]],ls=1.05)
bx=[(M,"PiFinder","exposes its solved\nposition as an\nINDI telescope",BLUE),
    (4.98,"Mount Bridge","couples that to\nany INDI mount\n(optional)",AMBER),
    (8.96,"Your rig","KStars / Ekos, the\nmount, the app —\nunchanged",RGBColor(0x8C,0xD8,0xB0))]
for x,t,d,c in bx:
    card(s,x,2.6,3.6,2.75,fill=DARK2,line=RGBColor(0x2C,0x3A,0x66))
    circle(s,x+0.34,2.94,0.18,c)
    text(s,x+0.66,2.82,2.9,0.4,[[(t,T_H2-4,True,WHITE)]])
    text(s,x+0.34,3.5,3.1,1.7,[[(ln,15,False,MOD)] for ln in d.split("\n")],ls=1.25)
for ax in (4.66,8.64):
    a=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(ax),Inches(3.55),Inches(0.42),Inches(0.55))
    a.fill.solid(); a.fill.fore_color.rgb=RGBColor(0x3A,0x4A,0x78); a.line.fill.background(); a.shadow.inherit=False
text(s,M,5.75,11.9,1.0,[[("PiFinder is patched in place, never forked. The mount side is generic INDI.",
    T_BODYSM-2,False,MOD)]],ls=1.3)
footer(s,"For dummies",dark=True)

# 6 SECTION 01
section("01","Part one","Project summary & focus")

# 7 WHAT PFSM IS
s=slide(LIGHT); head(s,"Scope","What PFSM is")
bullets(s,M,2.2,12.0,4.4,[
    ("A setup script ","— installs & patches PiFinder onto StellarMate OS."),
    ("diffs/*.diff patches ","— applied to a stock checkout via patch(1)."),
    ("Three INDI drivers ","— LX200, Mount Bridge, Simulator."),
    ("A web Control Center ","— stdlib-only, port 8765."),
],sz=T_BODY,gap=20)
footer(s,"Part 1 · Summary")

# 8 WHAT PFSM IS NOT
s=slide(LIGHT); head(s,"Scope","What PFSM is not")
bullets(s,M,2.2,12.0,4.4,[
    ("Not a fork of PiFinder ","— patches re-applied every update, kept minimal."),
    ("Not a fork of StellarMate ","— stock INDI, Web Manager, KStars."),
    ("Not a replacement for the INDI Control Panel ","— it stays for advanced work."),
    ("Not PiFinder-version-specific ","— talks only to the stable /api/*."),
],sz=T_BODY,gap=20,marker=RED)
footer(s,"Part 1 · Summary")

# 9 WHY INDI
s=slide(LIGHT); head(s,"The key decision","Why INDI")
text(s,M,2.4,12.0,1.6,[[("Both sides already speak it.",T_TITLE-4,True,NAVY)]],ls=1.05)
bullets(s,M,4.0,12.0,2.6,[
    "PiFinder exposes its position as an INDI telescope.",
    "The Mount Bridge couples that to any INDI mount — no per-mount code.",
],sz=T_BODY,gap=16)
footer(s,"Part 1 · Summary")

# 10 TESTED
s=slide(LIGHT); head(s,"Tested against","Versions & platforms")
rows=[("PiFinder 2.6.3  ·  SMOS 2.3.0","Pi 4  ✓   Pi 5  ✓   UTM x86  ✓ (install, CC, Mount Bridge vs. the Simulator)")]
text(s,M,2.5,12.0,1.0,[[("PiFinder 2.6.3   ·   StellarMate OS 2.3.0",T_H2,True,INK)]])
bullets(s,M,3.4,12.0,2.4,[
    ("Pi 4 ","— fully tested."),("Pi 5 ","— fully tested."),
    ("UTM x86 ","— install + Control Center + Mount Bridge vs. the Simulator (no real solving)."),
],sz=T_BODYSM,gap=14)
text(s,M,SH-1.3,12.0,0.9,[[("Pinned to a fixed release tag, not the upstream ",T_CAP+1,False,MUTE),
    ("release",T_CAP+1,True,INK),(" branch. The README table is the single source of truth.",T_CAP+1,False,MUTE)]],ls=1.3)
footer(s,"Part 1 · Summary")

# 11 SECTION 02
section("02","Part two","PFSM Control Center","Install, mode-switching, hardware checks and the Mount Bridge — one web page, no SSH.")

# 12 CC WHY
CCW=5.55; CCH=CCW/(2400/3030); CCX=SW-0.35-CCW; CCY=(SH-CCH)/2
s=slide(LIGHT); head(s,"Motivation","Why a Control Center")
bullets(s,M,2.35,6.5,4.2,[
    "Watch an install from a phone at the scope.",
    "Switch Fake Mode ⟷ the real service.",
    "Check camera / IMU / GPS are really detected.",
    "Reboot the Pi without SSH.",
],sz=T_BODY,gap=20)
pic(s,"cc_full.png",CCX,CCY,CCW,CCH)
footer(s,"Part 2 · Control Center")

# 13 CC STDLIB
s=slide(LIGHT); head(s,"Architecture","Stdlib only — by necessity")
text(s,M,2.4,6.5,3.0,[[("It bootstraps the very venv that PiFinder needs, "
    "so it can’t depend on one.",T_H2,True,INK)]],ls=1.2)
bullets(s,M,5.0,6.5,2.0,[
    "http.server on :8765, polled every 1–2 s.",
    "HTTP Basic auth against the system account (PAM).",
],sz=T_BODYSM,gap=14)
pic(s,"cc_full.png",CCX,CCY,CCW,CCH)
footer(s,"Part 2 · Control Center")

# 14 BADGES
s=slide(LIGHT); head(s,"Status language","Traffic-light badges")
rows=[("White","not applicable yet",RGBColor(0xC4,0xCE,0xDC)),
      ("Green","verified functional",GREEN),
      ("Yellow","degraded / estimate",AMBER),
      ("Red","broken or wrong",RED)]
yy=2.4
for t,d,c in rows:
    circle(s,M,yy+0.02,0.34,c)
    text(s,M+0.6,yy-0.05,11.0,0.5,[[(t+"  ",T_BODY,True,INK),(d,T_BODY,False,MUTE)]])
    yy+=0.78
text(s,M,SH-1.35,12.0,0.9,[[("The rule: verify real state, never “the process is alive”.",T_CAP+2,False,MUTE)]],ls=1.3)
footer(s,"Part 2 · Control Center")

# 15 CC COUPLING
CPH=6.0; CPW=CPH*(1328/1030); CPX=SW-0.35-CPW; CPY=(SH-CPH)/2
s=slide(LIGHT); head(s,"Operating it","Coupling, without the INDI panel")
bullets(s,M,2.55,4.5,4.2,[
    "Quick Actions — one-shots",
    "Coupling presets — one click",
    "Settings — set once, hidden",
],sz=T_BODYSM,gap=24)
pic(s,"cc_coupling.png",CPX,CPY,CPW,CPH)
footer(s,"Part 2 · Control Center")

# 16 SECTION 03
section("03","Part three","PF LX200 & PF Mount Bridge","PiFinder as a telescope, plus a generic bridge to any INDI mount.")

# 17 THREE BLOCKS
s=slide(LIGHT); head(s,"The INDI layer","Three building blocks")
blk=[("PiFinder LX200","INDI telescope. Reports the solved position; a GoTo becomes a push-to. Required.",BLUE),
     ("PiFinder Mount Bridge","Optional. Snoops PiFinder + the mount, couples them. Speaks only generic INDI.",AMBER),
     ("Your mount’s driver","Not ours — whatever your mount already uses. Only for a motorised mount.",RGBColor(0x8C,0xD8,0xB0))]
yy=2.2
for t,d,c in blk:
    circle(s,M,yy+0.06,0.32,c)
    text(s,M+0.6,yy-0.04,12.0,0.5,[[(t,T_H2-2,True,INK)]])
    text(s,M+0.6,yy+0.52,11.7,1.0,[[(d,T_BODYSM-2,False,MUTE)]],ls=1.22)
    yy+=1.62
footer(s,"Part 3 · LX200 & Mount Bridge")

# 18 LX200 DRIVER
s=slide(LIGHT); head(s,"The telescope driver","PiFinder LX200")
bullets(s,M,2.3,12.0,4.4,[
    ("Talks LX200 ","to PiFinder’s pos_server.py on TCP 4030."),
    ("Reports the freshly solved position ","every poll. A GoTo = a push-to target; nothing moves."),
    ("No Sync ","— PiFinder measures the truth every frame. Nothing to correct."),
],sz=T_BODY,gap=20)
footer(s,"Part 3 · LX200 & Mount Bridge")

# 19 COUPLING DIAL
s=slide(LIGHT); head(s,"One property, four positions","The Coupling Dial")
modes=[("Off","No coupling. Pure push-to.",MUTE),
       ("Verify / Alert only","Warn on drift. Never touches the mount.",BLUE),
       ("Auto-correct on drift","Past the threshold: Sync the mount to PiFinder.",AMBER),
       ("Goto-Forward","New target → real Goto, then settle & refine.",RGBColor(0x8C,0xD8,0xB0))]
yy=2.35
for t,d,c in modes:
    circle(s,M,yy+0.05,0.3,c)
    text(s,M+0.56,yy-0.05,4.7,0.9,[[(t,T_BODY-2,True,INK)]],ls=1.05)
    text(s,5.7,yy-0.05,7.1,0.9,[[(d,T_BODYSM-2,False,MUTE)]],ls=1.15)
    yy+=1.12
footer(s,"Part 3 · Auto-correct action: Sync (any mount) or Goto/Track")

# 20 GENERIC
s=slide(LIGHT); head(s,"Works with any mount","Generic INDI, nothing mount-specific")
text(s,M,2.35,12.0,1.4,[[("Only ",T_H2,False,INK),("EQUATORIAL_EOD_COORD",T_H2,True,NAVY),
    (" + ",T_H2,False,INK),("ON_COORD_SET",T_H2,True,NAVY),(" — never a mount command.",T_H2,False,INK)]],ls=1.2)
bullets(s,M,4.3,12.0,2.4,[
    ("Multi-Point Alignment ","(#191) — automated multi-star run."),
    ("Reposition detection ","— adopt an unexplained move, or revert."),
    ("Shadow Sync ","— mirror commands onto a non-driving device."),
],sz=T_BODYSM,gap=14)
footer(s,"Part 3 · LX200 & Mount Bridge")

# 21 SECTION 04
section("04","Part four","PF Simulator","Test the Mount Bridge with no real mount and no clear sky.")

# 22 SIM SETUP
s=slide(LIGHT); head(s,"Hardware-free testing","A mount side and a PiFinder side")
card(s,M,2.2,6.0,3.9)
text(s,M+0.45,2.5,5.0,0.5,[[("Mount side",T_H2,True,INK)]])
text(s,M+0.45,3.3,5.2,2.4,[[("The stock INDI Telescope Simulator — Sync it or GoTo it like a real mount.",
    T_BODYSM,False,MUTE)]],ls=1.28)
card(s,6.9,2.2,6.0,3.9)
text(s,7.35,2.5,5.0,0.5,[[("PiFinder side",T_H2,True,INK)]])
bullets(s,7.35,3.3,5.2,2.6,[
    ("Simulator ","— holds a settable RA/Dec."),
    ("Truth Injector ","— POSTs it to /api/fake_solve every ~2 s."),
],sz=T_BODYSM,gap=12)
footer(s,"Part 4 · Simulator")

# 23 SIM DETAILS
s=slide(LIGHT); head(s,"Behaviour & next","Simulator details")
bullets(s,M,2.3,12.0,3.6,[
    ("Sync == Goto ","— no motion to model; the RA/Dec just becomes the held value."),
    ("Inject every cycle ","— keeps PiFinder’s solve timestamp fresh."),
    ("FOLLOW_MOUNT_DEVICE ","(#239) — optionally dead-reckon along a mount’s slews."),
],sz=T_BODY,gap=18)
text(s,M,SH-1.5,12.0,1.0,[[("Next (v3.x): ",T_CAP+3,True,INK),
    ("render a GSC star field and run PiFinder’s real solver on it — issue #344.",T_CAP+3,False,MUTE)]],ls=1.3)
footer(s,"Part 4 · Simulator")

# 24 SECTION -> takeaway 1
section("→","Takeaway one · for PiFinder devs","Where PFSM could send PRs")

# 25 UPSTREAM
s=slide(LIGHT); head(s,"docs/upstream_patch_inventory.md","Patches that aren’t PFSM-specific")
bullets(s,M,2.2,12.2,4.2,[
    ("Test Mode toggle + status readback ","— a reliable API trigger (keypress sim was flaky)."),
    ("Camera-process resilience ","— fall back to the debug camera instead of crashing."),
    ("all_ips() ","— show every network address, not just the default route."),
    ("LX200 “no position yet” ","— return None, not a parseable fake coordinate."),
    ("Near-pole /api/fake_solve RA ","— filed as brickbots/PiFinder#645."),
],sz=T_BODYSM,gap=13)
text(s,M,SH-0.9,12.0,0.5,[[("Ready-to-file PR text: docs/upstream_pr_templates.md",T_CAP,False,MUTE)]])
footer(s,"Takeaway 1 · PiFinder PRs")

# 26 HOW TO HELP
s=slide(LIGHT); head(s,"Concretely","How to help — PiFinder")
bullets(s,M,2.3,12.0,4.0,[
    ("Pick one from the inventory ","— each is already triaged."),
    ("PR text is written ","in docs/upstream_pr_templates.md, in dependency order."),
    ("Mind the caveats ","— debug_solve needs an upstream regression fixed first."),
],sz=T_BODY,gap=18)
footer(s,"Takeaway 1 · PiFinder PRs")

# 27 SECTION -> takeaway 2
section("→","Takeaway two · for StellarMate devs","Extending the SMOS app for PFSM")

# 28 SMOS SCOPE
s=slide(LIGHT); head(s,"Basic usage in the app","What a thin integration needs")
bullets(s,M,2.2,12.0,4.0,[
    "Start / stop the selected INDI profile.",
    "Add / remove just the PiFinder drivers.",
    "Pick “the mount” from the loaded drivers.",
    "Three Coupling presets, in PiFinder’s words.",
],sz=T_BODY,gap=16)
text(s,M,SH-1.15,12.0,0.7,[[("The coupling logic is already framework-agnostic Python — a thin adapter ports it.",
    T_CAP+1,False,MUTE)]],ls=1.3)
footer(s,"Takeaway 2 · SMOS app")

# 29 SCREEN MAP
s=slide(LIGHT); head(s,"Mapping","CC feature → app screen")
pairs=[("Status badges","a read-only status card"),
       ("Coupling presets","a segmented control"),
       ("Quick Actions","buttons on that card")]
yy=2.5
for a,b in pairs:
    text(s,M,yy,5.4,0.6,[[(a,T_BODY-1,True,INK)]])
    ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(5.7),Inches(yy+0.06),Inches(0.5),Inches(0.42))
    ar.fill.solid(); ar.fill.fore_color.rgb=BLUE; ar.line.fill.background(); ar.shadow.inherit=False
    text(s,6.5,yy,6.2,0.6,[[(b,T_BODY-1,False,MUTE)]])
    yy+=1.15
text(s,M,SH-1.0,12.0,0.6,[[("APIs already exist: /api/mount_bridge_* , or the INDI properties directly.",T_CAP,False,MUTE)]])
footer(s,"Takeaway 2 · SMOS app")

# 30 ROADMAP
s=slide(DARK); stars(s,44,9)
text(s,M,0.7,11.6,0.8,[[("Roadmap & where to jump in",T_TITLE-2,True,WHITE)]])
card(s,M,1.95,6.0,4.3,fill=DARK2,line=RGBColor(0x2C,0x3A,0x66))
text(s,M+0.45,2.25,5.2,0.4,[[("v2.x — consolidate",T_H2-2,True,BLUE)]])
bullets(s,M+0.45,2.95,5.2,3.2,[
    "Harden & test what exists.",
    "Send selected patches upstream.",
    "A guiding watcher.",
],sz=T_BODYSM-2,gap=14,color=PAPER,marker=BLUE)
card(s,6.9,1.95,6.0,4.3,fill=DARK2,line=RGBColor(0x2C,0x3A,0x66))
text(s,7.35,2.25,5.2,0.4,[[("v3.x — integrate",T_H2-2,True,AMBER)]])
bullets(s,7.35,2.95,5.3,3.2,[
    "Deeper SMOS integration (#35, #36).",
    "PFSM actions in the PiFinder app (#43).",
    "A “real” GSC-solving simulator (#344).",
],sz=T_BODYSM-2,gap=14,color=PAPER,marker=AMBER)
text(s,M,6.55,12.0,0.4,[[("github.com/apos/PiFinder_Stellarmate",16,True,BLUE),
    ("   ·   GitHub Project #15   ·   CHANGELOG.md",14,False,MOD)]])

out=os.path.join(os.path.dirname(__file__),"PFSM_overview.pptx")
prs.save(out)
print("saved",out,"-",len(prs.slides._sldIdLst),"slides")
